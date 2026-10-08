"""Local CAS snapshot extraction and an explicit IBKR transaction adapter stub."""

import re
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import TYPE_CHECKING, Dict, Iterable, List, Optional, Sequence
from zipfile import ZipFile

from ingestion.base_parser import (
    BaseParser,
    ParsedHolding,
    ParsedTransaction,
    ParserNotImplementedError,
    ParsingError,
    UnsupportedStatementError,
)

if TYPE_CHECKING:
    from pdfplumber.page import Page


_MAX_BYTES = 20 * 1024 * 1024
_MAX_ROWS = 50_000
_ISIN = re.compile(r"\bIN[A-Z0-9]{10}\b")
_HEADERS = {
    "isin": {"isin", "isincode"},
    "security_name": {"security", "securityname", "nameofsecurity", "companyname", "company", "description", "schemename", "nameofthecompany", "nameofcompany", "nameofthecompanysecurity"},
    "quantity": {"quantity", "qty", "units", "balance", "closingbalance", "closingquantity", "currentbalance", "numberofunits", "noofshares", "noofunits", "numberofshares", "totalbalance"},
    "price": {"price", "marketprice", "unitprice", "nav", "rate", "pricers", "navrs"},
    "market_value": {"value", "valuation", "marketvalue", "currentvalue", "valueinrs", "valuers", "valuationrs", "marketvaluers", "currentvaluation"},
    "native_currency": {"currency", "nativecurrency"},
}


def _text(value: object) -> str:
    """Normalize whitespace in extracted cells without logging source data."""
    return " ".join(str(value).split()) if value is not None else ""


def _header(value: str) -> str:
    """Match headings independent of case, punctuation and line wrapping."""
    return re.sub(r"[^a-z]", "", value.lower())


def _columns(cells: Sequence[str]) -> Dict[str, int]:
    """Recognize only unambiguous holdings headings, never transaction headings."""
    columns: Dict[str, int] = {}
    for index, cell in enumerate(cells):
        for field, aliases in _HEADERS.items():
            if _header(cell) in aliases:
                if field in columns:
                    return {}
                columns[field] = index
    return columns if {"isin", "quantity"} <= columns.keys() else {}


def _transaction_header(cells: Sequence[str]) -> bool:
    """Identify transaction sections so their quantities never become snapshots."""
    headers = {_header(cell) for cell in cells}
    return bool(headers & {"transactiondate", "tradedate", "transactiontype", "credit", "debit", "narration"}) or (
        "date" in headers and bool(headers & {"isin", "isincode"})
    )


def _number(value: str, *, required: bool = False) -> Optional[Decimal]:
    """Parse plain or Indian-grouped nonnegative decimals without float math."""
    if value in {"", "-", "--", "N/A"} and not required:
        return None
    cleaned = re.sub(r"^(?:INR|Rs\.?)\s*", "", value, flags=re.IGNORECASE).strip()
    if not re.fullmatch(r"(?:\d+|\d{1,3}(?:,\d{3})+|\d{1,2}(?:,\d{2})*,\d{3})(?:\.\d+)?", cleaned):
        raise ParsingError("A holdings row contains an invalid or missing numeric value.")
    try:
        number = Decimal(cleaned.replace(",", ""))
    except InvalidOperation:
        raise ParsingError("A holdings row contains an invalid numeric value.") from None
    if not number.is_finite() or number < 0:
        raise ParsingError("A holdings row contains an invalid numeric value.")
    return number


def _parse_rows(rows: Iterable[Sequence[object]], source: str) -> List[ParsedHolding]:
    """Extract rows under recognized headers; reject ambiguous ISIN-bearing rows."""
    holdings: List[ParsedHolding] = []
    columns: Dict[str, int] = {}
    previous: List[str] = []
    transaction_section = False
    for row_index, row in enumerate(rows):
        if row_index >= _MAX_ROWS:
            raise ParsingError("The statement exceeds the supported row limit.")
        cells = [_text(cell) for cell in row]
        if _transaction_header(cells):
            columns = {}
            previous = []
            transaction_section = True
            continue
        recognized = _columns(cells)
        if not recognized and previous and len(previous) == len(cells):
            recognized = _columns([f"{above} {below}" for above, below in zip(previous, cells)])
        if recognized:
            columns = recognized
            previous = []
            transaction_section = False
            continue
        if any(_header(cell) in {"isin", "isincode"} for cell in cells):
            columns = {}
        previous = cells
        if transaction_section:
            continue
        matches = _ISIN.findall(" ".join(cells))
        if not matches:
            if columns and columns["isin"] < len(cells):
                identifier = cells[columns["isin"]]
                quantity_text = cells[columns["quantity"]] if columns["quantity"] < len(cells) else ""
                total_row = any(_header(cell) in {"total", "subtotal", "grandtotal"} for cell in cells)
                if not total_row and (identifier or quantity_text):
                    raise ParsingError("A holdings row is missing a valid ISIN. No partial preview was accepted.")
            continue
        if not columns or len(matches) != 1 or max(columns.values()) >= len(cells):
            raise ParsingError("An ISIN row has no unambiguous holdings columns. This layout is not supported.")
        values = {field: cells[index] for field, index in columns.items()}
        if not _ISIN.fullmatch(values["isin"]):
            raise ParsingError("A holdings row has an invalid ISIN column.")
        quantity = _number(values["quantity"], required=True)
        if quantity is None:
            raise ParsingError("A holdings row is missing its quantity.")
        currency = values.get("native_currency", "INR").upper()
        if currency != "INR":
            raise ParsingError("Only INR-denominated CAS holdings are supported.")
        holdings.append(ParsedHolding(
            isin=values["isin"], security_name=values.get("security_name", ""),
            quantity=quantity, price=_number(values.get("price", "")),
            market_value=_number(values.get("market_value", "")),
            native_currency=currency, source=source,
        ))
    return holdings


def _validate_file(file_path: Path, suffix: str) -> None:
    """Validate only file metadata, without reading or logging private contents."""
    try:
        valid = file_path.suffix.lower() == suffix and file_path.is_file()
    except OSError:
        raise UnsupportedStatementError("The statement file is inaccessible.") from None
    if not valid:
        raise UnsupportedStatementError("The statement file is missing or has an unsupported format.")


def _pdf_text_rows(page: "Page") -> List[Sequence[object]]:
    """Use word coordinates to separate columns without relying on rendered spaces."""
    rows: List[Sequence[object]] = []
    cells: List[str] = []
    line_top: Optional[float] = None
    previous_right = 0.0
    previous_width = 0.0
    for word in sorted(page.extract_words(x_tolerance=1), key=lambda item: (item["top"], item["x0"])):
        top = float(word["top"])
        left = float(word["x0"])
        right = float(word["x1"])
        text = str(word["text"])
        if line_top is None or abs(top - line_top) > 3:
            if cells:
                rows.append(cells)
            cells = [text]
            line_top = top
        elif left - previous_right > max(8, previous_width * 1.8):
            cells.append(text)
        else:
            cells[-1] += " " + text
        previous_right = right
        previous_width = (right - left) / max(1, len(text))
    if cells:
        rows.append(cells)
    return rows


class CdslCasParser(BaseParser[ParsedHolding]):
    """Read tabular CAS snapshots locally, preserving statement valuation fields.

    Supports text-based PDF tables and XLSX holdings sheets with ISIN and closing
    quantity headings. Scanned PDFs and unknown layouts are rejected, not guessed.
    Uploaded bytes and PDF passwords are never written to disk by this parser.
    """

    def parse_file(self, file_path: Path) -> List[ParsedHolding]:
        """Read a bounded local PDF/XLSX; use parse_bytes for password-protected PDFs."""
        suffix = file_path.suffix.lower()
        if suffix not in {".pdf", ".xlsx"}:
            raise UnsupportedStatementError("Select a PDF or XLSX CAS statement.")
        _validate_file(file_path, suffix)
        try:
            with file_path.open("rb") as document:
                content = document.read(_MAX_BYTES + 1)
        except OSError:
            raise ParsingError("The statement could not be read.") from None
        return self.parse_bytes(content, suffix)

    def parse_bytes(self, content: bytes, suffix: str, password: Optional[str] = None) -> List[ParsedHolding]:
        """Parse an in-memory upload; sanitize third-party exceptions at the boundary."""
        if not content or len(content) > _MAX_BYTES:
            raise ParsingError("Select a nonempty statement no larger than 20 MiB.")
        try:
            if suffix.lower() == ".pdf":
                holdings = self._pdf(content, password)
            elif suffix.lower() == ".xlsx":
                holdings = self._xlsx(content)
            else:
                raise UnsupportedStatementError("Select a PDF or XLSX CAS statement.")
        except ParsingError:
            raise
        except Exception:
            raise ParsingError("The statement could not be parsed. Check its format and PDF password.") from None
        if not holdings:
            raise ParsingError("No supported holdings table found. Scanned PDFs require OCR; unknown layouts are not imported.")
        return holdings

    @staticmethod
    def _xlsx(content: bytes) -> List[ParsedHolding]:
        """Read worksheet values without evaluating formulas or external links."""
        from openpyxl import load_workbook

        with ZipFile(BytesIO(content)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 100 * 1024 * 1024:
                raise ParsingError("The expanded workbook exceeds the supported size limit.")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            if len(workbook.worksheets) > 50:
                raise ParsingError("The workbook exceeds the supported sheet limit.")
            holdings: List[ParsedHolding] = []
            for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
                if (sheet.max_row or 0) > _MAX_ROWS or (sheet.max_column or 0) > 100:
                    raise ParsingError("The worksheet exceeds the supported dimensions.")
                holdings.extend(_parse_rows(sheet.iter_rows(values_only=True), f"Sheet {sheet_index}"))
            return holdings
        finally:
            workbook.close()

    @staticmethod
    def _pdf(content: bytes, password: Optional[str]) -> List[ParsedHolding]:
        """Decrypt in memory and extract ruled or whitespace-aligned text tables."""
        import pdfplumber
        from pypdf import PdfReader, PdfWriter

        if not content.startswith(b"%PDF-"):
            raise ParsingError("The selected file is not a PDF document.")
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted and not reader.decrypt(password or ""):
            raise ParsingError("The PDF password is missing or incorrect.")
        if len(reader.pages) > 200:
            raise ParsingError("The PDF exceeds the supported page limit.")
        decrypted = BytesIO()
        writer = PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        writer.write(decrypted)
        decrypted.seek(0)
        holdings: List[ParsedHolding] = []
        with pdfplumber.open(decrypted) as document:
            for page_index, page in enumerate(document.pages, start=1):
                source = f"Page {page_index}"
                tables = page.extract_tables()
                rows: List[Sequence[object]] = [row for table in tables for row in table]
                if not any(_columns([_text(cell) for cell in row]) for row in rows):
                    rows = _pdf_text_rows(page)
                holdings.extend(_parse_rows(rows, source))
        return holdings


class IbkrFlexParser(BaseParser[ParsedTransaction]):
    """Future offline IBKR/Paasa Flex XML adapter; no tokens or API calls."""

    def parse_file(self, file_path: Path) -> List[ParsedTransaction]:
        """Validate an XML path and raise ParserNotImplementedError for this stub."""
        _validate_file(file_path, ".xml")
        raise ParserNotImplementedError("IBKR Flex parsing is not implemented yet.")